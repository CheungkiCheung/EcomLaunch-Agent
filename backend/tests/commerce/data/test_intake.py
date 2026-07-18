"""Deterministic contracts for safe heterogeneous data intake."""

from __future__ import annotations

import json
import stat
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.commerce.data.intake import DataIntakeError, DataIntakeService, FileFormat
from app.commerce.domain.ids import WorkspaceId


def _service(tmp_path: Path) -> DataIntakeService:
    return DataIntakeService(storage_root=tmp_path / "commerce-storage")


def test_csv_intake_records_hash_encoding_table_and_read_only_raw_file(tmp_path: Path):
    source = tmp_path / "orders.csv"
    source.write_text("order_id,seller_id\norder-1,seller-1\n", encoding="utf-8")

    manifest = _service(tmp_path).ingest(WorkspaceId.new(), (source,))

    assert len(manifest.files) == 1
    assert len(manifest.tables) == 1
    file = manifest.files[0]
    table = manifest.tables[0]
    assert file.original_name == "orders.csv"
    assert file.format is FileFormat.CSV
    assert file.encoding == "utf-8"
    assert len(file.sha256) == 64
    assert table.table_name == "orders"
    assert table.source_file_id == file.id

    stored_path = tmp_path / "commerce-storage" / manifest.storage_relative_path / file.stored_relative_path
    assert stored_path.read_text(encoding="utf-8").startswith("order_id")
    assert stored_path.stat().st_mode & stat.S_IWUSR == 0
    assert (stored_path.parents[1] / "manifest.json").is_file()


def test_json_object_of_tables_creates_multiple_table_manifests(tmp_path: Path):
    source = tmp_path / "bundle.json"
    source.write_text(
        json.dumps(
            {
                "orders": [{"order_id": "order-1"}],
                "reviews": [{"order_id": "order-1", "score": 5}],
            }
        ),
        encoding="utf-8",
    )

    manifest = _service(tmp_path).ingest(WorkspaceId.new(), (source,))

    assert {table.table_name for table in manifest.tables} == {"orders", "reviews"}


def test_excel_intake_records_each_sheet_as_a_table(tmp_path: Path):
    source = tmp_path / "merchant.xlsx"
    workbook = Workbook()
    orders = workbook.active
    orders.title = "Orders"
    orders.append(("order_id", "seller_id"))
    orders.append(("order-1", "seller-1"))
    reviews = workbook.create_sheet("Reviews")
    reviews.append(("order_id", "score"))
    reviews.append(("order-1", 5))
    workbook.save(source)

    manifest = _service(tmp_path).ingest(WorkspaceId.new(), (source,))

    assert {table.table_name for table in manifest.tables} == {"orders", "reviews"}
    assert {table.sheet_name for table in manifest.tables} == {"Orders", "Reviews"}


def test_zip_intake_safely_expands_supported_table_members(tmp_path: Path):
    source = tmp_path / "bundle.zip"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("exports/orders.csv", "order_id\norder-1\n")
        archive.writestr("exports/reviews.jsonl", '{"order_id":"order-1","score":5}\n')

    manifest = _service(tmp_path).ingest(WorkspaceId.new(), (source,))

    assert manifest.files[0].format is FileFormat.ZIP
    assert {table.table_name for table in manifest.tables} == {"orders", "reviews"}
    member_files = [file for file in manifest.files if file.archive_member]
    assert {file.archive_member for file in member_files} == {
        "exports/orders.csv",
        "exports/reviews.jsonl",
    }
    for file in member_files:
        stored_path = tmp_path / "commerce-storage" / manifest.storage_relative_path / file.stored_relative_path
        assert stored_path.stat().st_mode & stat.S_IWUSR == 0


@pytest.mark.parametrize(
    "member_name",
    (
        "../outside.csv",
        "/absolute.csv",
        "safe/../../outside.csv",
    ),
)
def test_zip_path_traversal_is_rejected(tmp_path: Path, member_name: str):
    source = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr(member_name, "id\n1\n")

    with pytest.raises(DataIntakeError, match="unsafe ZIP member"):
        _service(tmp_path).ingest(WorkspaceId.new(), (source,))

    assert not (tmp_path / "outside.csv").exists()


def test_zip_duplicate_members_are_rejected(tmp_path: Path):
    source = tmp_path / "duplicate.zip"
    with pytest.warns(UserWarning, match="Duplicate name"):
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("orders.csv", "id\n1\n")
            archive.writestr("orders.csv", "id\n2\n")

    with pytest.raises(DataIntakeError, match="duplicate ZIP member"):
        _service(tmp_path).ingest(WorkspaceId.new(), (source,))


def test_symbolic_link_input_is_rejected(tmp_path: Path):
    target = tmp_path / "orders.csv"
    target.write_text("id\n1\n", encoding="utf-8")
    link = tmp_path / "orders-link.csv"
    link.symlink_to(target)

    with pytest.raises(DataIntakeError, match="symbolic links"):
        _service(tmp_path).ingest(WorkspaceId.new(), (link,))


def test_unsupported_file_type_is_rejected(tmp_path: Path):
    source = tmp_path / "notes.txt"
    source.write_text("not a supported table", encoding="utf-8")

    with pytest.raises(DataIntakeError, match="Unsupported input format"):
        _service(tmp_path).ingest(WorkspaceId.new(), (source,))
